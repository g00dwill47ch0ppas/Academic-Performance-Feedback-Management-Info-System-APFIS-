from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "manuals"
STILLS = ROOT / "product-film" / "output" / "stills"
OUT.mkdir(parents=True, exist_ok=True)

PURPLE = "4C237F"
VIOLET = "7C4DCC"
PALE = "F2EDF8"
INK = "20202A"
MUTED = "626273"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=INK, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, "FFFFFF", 9)
        shade(t.rows[0].cells[i], PURPLE)
    header_pr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, False, INK, 9)
            if ridx % 2 == 1:
                shade(cells[i], "F7F5FA")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p

def number(doc, text):
    if not hasattr(doc, "_manual_step"):
        doc._manual_step = 0
    doc._manual_step += 1
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.add_run(f"{doc._manual_step}. {text}")
    return p

def reset_steps(doc):
    doc._manual_step = 0

def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.18)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(PURPLE)
    return p

def setup_doc(title, subtitle, audience):
    d = Document()
    sec = d.sections[0]
    sec.top_margin = Inches(.72)
    sec.bottom_margin = Inches(.7)
    sec.left_margin = Inches(.78)
    sec.right_margin = Inches(.78)
    styles = d.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = RGBColor.from_string(INK)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 28), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)]:
        s = styles[name]
        s.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(PURPLE if name != "Title" else INK)
        s.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        s.paragraph_format.space_after = Pt(5)
        s.paragraph_format.keep_with_next = True
        if name == "Title":
            ppr = s._element.get_or_add_pPr()
            borders = ppr.find(qn("w:pBdr"))
            if borders is not None:
                ppr.remove(borders)
    subtitle_ppr = styles["Subtitle"]._element.get_or_add_pPr()
    subtitle_borders = subtitle_ppr.find(qn("w:pBdr"))
    if subtitle_borders is not None:
        subtitle_ppr.remove(subtitle_borders)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = header.add_run("APFIS  |  " + audience.upper())
    rr.font.name = "Aptos"
    rr.font.size = Pt(8)
    rr.font.color.rgb = RGBColor.from_string(MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Academic Performance Feedback Management Information System  •  29 October 2026")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string(MUTED)
    d.add_paragraph("ACADEMIC PERFORMANCE FEEDBACK MANAGEMENT INFORMATION SYSTEM", style="Subtitle")
    p = d.add_paragraph(style="Title")
    p.add_run(title)
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(15)
    r = p.add_run(subtitle)
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(VIOLET)
    for title_p in d.paragraphs[-2:]:
        ppr = title_p._p.get_or_add_pPr()
        borders = ppr.find(qn("w:pBdr"))
        if borders is not None:
            ppr.remove(borders)
    table(d, ["Document details", "Value"], [
        ("Prepared for", "APFIS project handover and university project review"),
        ("Authors", "Gosego Sebaetse and Thamsanqa Ndlakuse"),
        ("Version", "1.0"),
        ("Prepared", "29 October 2026"),
        ("Audience", audience),
    ], [1.45, 5.35])
    d.add_heading("About this manual", level=1)
    d.add_paragraph(subtitle + " This guide describes the APFIS implementation in the repository snapshot reviewed for this document. The example manuals supplied with the request informed the section structure only; their AACMS features and claims are not part of APFIS.")
    d.add_paragraph("Implementation note: APFIS is a prototype for local demonstration and academic review. It is not connected to an institutional learning management system, student information system, email service, or production database.")
    d.add_page_break()
    return d

def add_contents(d, headings):
    d.add_heading("Contents", level=1)
    for h in headings:
        bullet(d, h)
    d.add_page_break()

def add_figure(d, filename, caption, width=6.6):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(STILLS / filename), width=Inches(width))
    cap = d.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(9)
    cap.paragraph_format.keep_together = True
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)

def technical_manual():
    d = setup_doc("Developer Technical Manual", "Architecture, local setup, data flow and implementation notes", "Developers, maintainers and technical reviewers")
    add_contents(d, [
        "1. Purpose and implementation scope", "2. System architecture", "3. Technology and prerequisites", "4. Repository map", "5. Local setup and operation", "6. Data model and privacy boundaries", "7. Calculations and decision logic", "8. Routes and user interface", "9. Validation, testing and maintenance", "10. Limitations and future integration", "Appendix A. Common commands", "Appendix B. Source files to inspect first", "Appendix C. Glossary",
    ])
    d.add_heading("1. Purpose and implementation scope", level=1)
    d.add_paragraph("This manual helps a developer run APFIS locally, locate the main implementation areas, understand how the two applications handle data, and maintain the calculation workflows. It documents behaviour found in the current source code. It does not treat research goals, earlier specifications, or the example AACMS manuals as evidence that an unimplemented feature exists.")
    d.add_paragraph("APFIS is a suite with a landing page and two independent Flask applications: a lecturer Decision Support System (DSS) and a Student Performance Assistant. The lecturer tool analyses a cohort and assessment plan. The student tool lets one student model a module and explore marks needed in remaining assessments.")
    d.add_heading("2. System architecture", level=1)
    table(d, ["Component", "Implementation", "Data boundary"], [
        ("Landing page", "Small Flask server serving gateway/index.html on port 5000", "Links to the two apps; it does not proxy or exchange academic records."),
        ("Lecturer DSS", "Flask app factory, routes, templates, Python calculation services", "One process-level in-memory store; reset when that process restarts."),
        ("Student Performance Assistant", "Flask serves Jinja templates and static JavaScript; browser computes workflow results", "Student model and intermediate calculations use browser localStorage."),
    ], [1.2, 2.5, 3.1])
    d.add_paragraph("The separate applications are intentional. They each contain their own backend package and app/config modules. The landing page links to each local service; no API, shared database, common session, or data synchronisation between lecturer and student views is implemented.")
    code(d, "Browser → gateway :5000 → lecturer DSS :5001\n                          ↘ Student Performance Assistant :5002\nLecturer browser ↔ lecturer Flask process memory\nStudent browser localStorage ↔ student workflow JavaScript")
    d.add_heading("3. Technology and prerequisites", level=1)
    table(d, ["Area", "Observed technology"], [
        ("Runtime", "Python 3; Flask applications"),
        ("Lecturer calculations", "NumPy, pandas, SciPy optimisation; PuLP is listed among project dependencies"),
        ("Student interface", "Flask/Jinja templates, Bootstrap styling, browser JavaScript and localStorage"),
        ("Persistence", "No database is configured in the inspected apps; lecturer store is in memory and student data is browser-local"),
        ("External service", "No LMS integration, email/SMS delivery, or remote student-data source is present"),
    ], [1.5, 5.2])
    d.add_paragraph("Use a current Python 3 installation and a modern browser. Install the repository-level requirements into an isolated virtual environment. Verify package versions in the requirements files when reproducing a specific build.")
    d.add_heading("4. Repository map", level=1)
    table(d, ["Path", "Responsibility"], [
        ("gateway/", "Landing page server and HTML entry point."),
        ("apps/lecturer/app.py, config.py", "Lecturer Flask entry point and configuration."),
        ("apps/lecturer/backend/routes/", "Module, student, cohort, home/dashboard and tutorial routes."),
        ("apps/lecturer/backend/algorithms/", "Feasible bounds, target participation plan and cohort weight solver."),
        ("apps/lecturer/backend/services/participation.py", "Shared best-N-of-M participation calculation and ranking rules."),
        ("apps/lecturer/backend/data/", "In-memory store, data loader and simulated sample CSV."),
        ("apps/lecturer/frontend/templates/ and static/", "Lecturer pages, styles and browser interactions."),
        ("apps/student/app.py and backend/", "Student Flask app and page routes."),
        ("apps/student/frontend/templates/", "Welcome, returning student and four workflow phase screens."),
        ("apps/student/frontend/static/js/app.js", "Student-side model, profile, improvement-option and summary calculations."),
        ("apps/student/frontend/static/js/returning.js", "CSV import and model reconstruction for a returning student."),
        ("product-film/output/stills/", "Approved film stills used as illustrative interface references in the user manual."),
    ], [2.6, 4.1])
    d.add_heading("5. Local setup and operation", level=1)
    d.add_heading("5.1 Create an environment and install packages", level=2)
    code(d, "cd <repository-root>\npython -m venv .venv\n.venv\\Scripts\\Activate.ps1\npython -m pip install -r requirements.txt")
    d.add_heading("5.2 Start the three local processes", level=2)
    d.add_paragraph("Open three PowerShell windows from the repository root. Activate the same virtual environment in each window.")
    code(d, "# Window 1\ncd apps\\lecturer\nflask --app app run --port 5001\n\n# Window 2\ncd apps\\student\nflask --app app run --port 5002\n\n# Window 3\ncd gateway\npython gateway.py")
    d.add_paragraph("Open http://localhost:5000. The landing page checks the local services and links to the lecturer and student tools. If started directly, the lecturer app is at port 5001 and the student app at port 5002.")
    d.add_heading("5.3 Configuration and development safety", level=2)
    d.add_paragraph("The lecturer Config reads SECRET_KEY, FLASK_DEBUG, DEFAULT_PASS_THRESHOLD and DEFAULT_TARGET_CLASS_AVERAGE. Its current defaults include a development secret and debug mode. Do not expose this configuration as a production service. Set a private secret and disable debug when running outside a trusted local development session. The student Flask app is also configured for local use and does not provide account authentication.")
    d.add_heading("6. Data model and privacy boundaries", level=1)
    d.add_heading("6.1 Lecturer application", level=2)
    d.add_paragraph("The lecturer DataStore holds modules, assessment plans, students and marks in process memory. The application factory seeds a simulated CS101 module from apps/lecturer/backend/data/sample_data.csv when the store is empty. Module data is lost when the lecturer server restarts. CSV import and export are available. The importer includes validation intended to reject personally identifiable fields; confirm the current accepted headers and validation rules in backend/data/loader.py before preparing a file.")
    d.add_heading("6.2 Student application", level=2)
    d.add_paragraph("The Student Performance Assistant stores studentModel and phase calculation objects in browser localStorage. Browser storage is specific to the browser profile and origin. Clearing site data or changing browser/profile can remove or isolate the model. The CSV summary export allows a user to save a copy and the returning-student workflow imports that application-specific file format. Avoid entering real personal records in demonstration environments.")
    d.add_heading("6.3 Separation and data handling", level=2)
    bullet(d, "The landing page only navigates; the apps do not share records or send marks between each other.")
    bullet(d, "The lecturer app has no external database in this build; its server memory is not durable storage.")
    bullet(d, "The student browser stores user-entered details locally; this is not an institutional account store or encrypted remote service.")
    bullet(d, "No login, role-based access control, email, push notification, or LMS event ingestion is implemented in the inspected source.")
    d.add_heading("7. Calculations and decision logic", level=1)
    d.add_heading("7.1 Lecturer participation and risk flag", level=2)
    d.add_paragraph("Assessment weights are applied through a shared participation service. For a type with M items and N counted items, the completed items are ranked by percentage mark; the best N count, with earliest plan order breaking ties. Each counted slot contributes the type weight divided by N. A blank mark is treated as not yet assessed, not as zero.")
    d.add_paragraph("The current participation mark is the accrued weighted contribution. The feasible range runs from that accrued amount to the best achievable value if remaining items receive full marks and can displace lower completed items. A student is marked at risk only when their maximum feasible final mark is strictly below the module pass threshold. The flag is therefore a rule-based best-case test, not a trained prediction model or probability.")
    d.add_heading("7.2 Individual target plan", level=2)
    d.add_paragraph("After showing the feasible region, the lecturer selects an absolute target inside it. The planner uses monotone bisection to find the minimum uniform percentage for the remaining items that reaches the target. It also evaluates whether any remaining items can be skipped under best-N-of-M and calculates the required percentage for the remaining active items. Values are rounded for display only.")
    d.add_heading("7.3 Cohort weight planning", level=2)
    d.add_paragraph("The lecturer supplies a target cohort average and minimum/maximum weight ranges for each assessment type. The solver uses SciPy's constrained nonlinear optimisation. Ranges must have positive minimums, each minimum must not exceed its maximum, and the aggregate ranges must allow weights to sum to 100%. An assessment weight cannot be zero; a 1% floor is enforced. The interface reports the solver outcome and achieved average. This is an optimisation scenario, not a guarantee that changing weights will cause a target average in a real cohort.")
    d.add_paragraph("Implementation notes in the source mark some methodological differences from the referenced research paper as open items. In particular, the best-N-of-M extension and the treatment of not-yet-completed marks in cohort optimisation are project-specific. Review apps/lecturer/docs/technical_specification.md and algorithm comments before making research claims.")
    d.add_heading("7.4 Student-side profile and options", level=2)
    d.add_paragraph("The student enters an assessment framework and existing marks. Browser JavaScript calculates an academic profile and improvement options from the stored model, then Phase 4 displays a read-only summary and CSV export. The result is a personal planning aid based on entered information; it is not an official grade, lecturer recommendation, or direct student-facing message from the lecturer DSS.")
    d.add_heading("8. Routes and user interface", level=1)
    table(d, ["App area", "Implemented pages or actions"], [
        ("Gateway", "Landing page at /; provides links/status for local services."),
        ("Lecturer", "Overview/dashboard, at-risk list, student list/detail/add/edit, CSV imports, modules, assessment plan, cohort planning, settings and tutorial."),
        ("Student", "Welcome/new student, returning student CSV import, Phase 1 assessment plan, Phase 2 profile, Phase 3 improvement options, Phase 4 summary."),
    ], [1.2, 5.5])
    d.add_paragraph("The lecturer interface uses a purple accent and Bootstrap-like card/table layouts; the student assistant uses purple branding, Poppins/Inter typography, a responsive card layout, and a light/dark theme preference. The user manual includes stills from the real screen structure for orientation.")
    d.add_heading("9. Validation, testing and maintenance", level=1)
    d.add_paragraph("The lecturer project contains unit tests under apps/lecturer/tests for participation rules, bounds, target plans, cohort planning, rounding, imports/exports, the store and tutorial behavior. Run the test suite from apps/lecturer using the project's configured Python environment before changing calculation logic. No performance, accessibility, security, or institutional integration certification is implied by these tests.")
    d.add_paragraph("When changing calculation rules, keep participation.py as the shared rule implementation and preserve route-level validation. Add regression cases for blank marks, equal-percentage ties, best-N drops, unreachable targets and invalid cohort ranges. Use synthetic records in screenshots and test fixtures.")
    d.add_heading("10. Limitations and future integration", level=1)
    for s in [
        "No LMS logins, attendance events, submission timestamps or live grades are ingested automatically.",
        "The lecturer risk flag is threshold-based and does not learn from historical student outcomes.",
        "No student notification is sent by the lecturer tool; no email, push alert or messaging integration exists.",
        "The lecturer data store is volatile and the apps have no production authentication or user isolation.",
        "The two tools are intentionally separate in this prototype; privacy separation should remain explicit in any future integration design.",
        "The supplied lecturer sample cohort is simulated. Never interpret sample dashboard values as research results or institutional performance.",
    ]:
        bullet(d, s)
    d.add_heading("Appendix A. Common commands", level=1)
    table(d, ["Task", "Command"], [
        ("Install dependencies", "python -m pip install -r requirements.txt"),
        ("Run lecturer app", "cd apps/lecturer; flask --app app run --port 5001"),
        ("Run student app", "cd apps/student; flask --app app run --port 5002"),
        ("Run gateway", "cd gateway; python gateway.py"),
        ("Inspect working tree", "git status --short"),
    ], [1.65, 5.0])
    d.add_heading("Appendix B. Source files to inspect first", level=1)
    table(d, ["Concern", "Starting point"], [
        ("Lecturer app startup", "apps/lecturer/backend/__init__.py; apps/lecturer/config.py"),
        ("Participation/risk logic", "apps/lecturer/backend/services/participation.py; algorithms/lp_bounds.py"),
        ("Individual plan", "apps/lecturer/backend/algorithms/participation_plan.py"),
        ("Cohort optimisation", "apps/lecturer/backend/algorithms/nlp_weights.py; routes/cohort.py"),
        ("Import and state", "apps/lecturer/backend/data/loader.py; data/store.py"),
        ("Student workflow", "apps/student/frontend/static/js/app.js; templates/phase*.html"),
    ], [1.8, 4.9])
    d.add_heading("Appendix C. Glossary", level=1)
    table(d, ["Term", "Meaning in APFIS"], [
        ("Participation mark", "Weighted contribution already accrued across counted assessment slots."),
        ("Standing", "Average across currently counted work, expressed as a percentage."),
        ("Feasible region", "Minimum and maximum final mark reachable under the assessment plan."),
        ("Best N of M", "Only the highest N percentage results among M items of a type count."),
        ("At risk", "Best-case final mark is below the configured pass threshold."),
        ("Prescriptive planning", "The system explores what marks or weight ranges could reach a selected target; it does not deliver a message to a student."),
    ], [1.8, 4.9])
    d.save(OUT / "APFIS_Developers_Technical_Manual.docx")

def user_manual():
    d = setup_doc("User Manual", "A practical guide for students and lecturers", "Students, lecturers and project-day reviewers")
    add_contents(d, [
        "1. About APFIS", "2. Start the applications", "3. Student Performance Assistant", "4. Lecturer Decision Support System", "5. Understand the results", "6. Privacy and responsible use", "7. Troubleshooting", "8. Quick workflow reference",
    ])
    d.add_heading("1. About APFIS", level=1)
    d.add_paragraph("APFIS provides two separate tools for academic performance planning. The Student Performance Assistant helps a student enter assessment details and marks, review a calculated profile, and explore improvement options. The lecturer Decision Support System helps a lecturer review a module cohort, identify students whose best-case mark is below the pass threshold, plan an individual target, and explore assessment-weight ranges.")
    d.add_paragraph("The tools do not share student records. The student assistant does not receive a lecturer alert or message, and the lecturer DSS does not send recommendations to a student's device. A lecturer and student can discuss the separate results outside the application.")
    d.add_heading("2. Start the applications", level=1)
    d.add_paragraph("For the local demonstration, start the student app at port 5002, the lecturer app at port 5001, and the landing page at port 5000. Open http://localhost:5000 in a modern browser and choose the tool you need. If you are given a direct local link, use the Student Performance Assistant on port 5002 or the Lecturer Decision Support System on port 5001.")
    d.add_paragraph("The lecturer app starts with a simulated CS101 sample module. It is provided so the dashboard can be explored immediately. Do not treat those sample values as official or real student results.")
    d.add_heading("3. Student Performance Assistant", level=1)
    d.add_heading("3.1 Begin a new student model", level=2)
    reset_steps(d)
    number(d, "On the welcome screen, choose the new-student path.")
    number(d, "In Phase 1, enter the student details and module, then define assessment categories, item counts and weights. Confirm the assessment framework before continuing.")
    number(d, "Select Continue to Phase 2. The assessment categories expand into individual assessment rows.")
    add_figure(d, "scene-03-setup.png", "Figure 1. Student assistant Phase 1 assessment setup. Use fictional values for a demonstration.")
    d.add_heading("3.2 Enter marks and review the profile", level=2)
    reset_steps(d)
    number(d, "In Phase 2, enter marks for assessments that have been completed. Leave future assessments blank; a blank is not the same as a mark of zero.")
    number(d, "Select Generate Model to calculate the profile from the assessment framework and marks you entered.")
    number(d, "Review current participation, academic standing, completed work and remaining assessments. Check that the module and marks are correct before using the estimate.")
    add_figure(d, "scene-04-profile.png", "Figure 2. Phase 2 profile summarises the entered assessment results.")
    d.add_heading("3.3 Explore improvement options", level=2)
    reset_steps(d)
    number(d, "Continue to Phase 3 and select a target/improvement option supported by the screen.")
    number(d, "Review the feasible option cards and any required result shown for remaining assessment items.")
    number(d, "Treat the values as planning estimates based on your inputs. They are not an official final grade or a guarantee of a future result.")
    add_figure(d, "scene-05-scenarios.png", "Figure 3. Phase 3 displays improvement options calculated from the entered marks.")
    d.add_heading("3.4 Review and export the summary", level=2)
    reset_steps(d)
    number(d, "Open Phase 4 to review the student information, assessment framework, marks and displayed profile/options.")
    number(d, "Use the CSV export control if you want to save the assessment information for your own use.")
    number(d, "To continue later, retain the downloaded APFIS CSV. On the welcome screen choose the returning-student option, select that file and import it.")
    d.add_paragraph("The returning workflow expects the CSV structure created by this application. Do not edit the file structure manually. The CSV contains student details and assessment framework/marks; it is not a system backup of every calculated value.")
    add_figure(d, "scene-06-summary.png", "Figure 4. Phase 4 read-only student performance summary.")
    d.add_heading("4. Lecturer Decision Support System", level=1)
    d.add_heading("4.1 Open the module dashboard", level=2)
    reset_steps(d)
    number(d, "Open the lecturer DSS. The sample module dashboard displays cohort-level participation, dropped marks, at-risk count and class size.")
    number(d, "Use the navigation panel to open the student list, module/assessment plan, at-risk view, settings or tutorial.")
    number(d, "Use the module controls to create a module, import a supported CSV file, or switch modules. Follow the page validation messages.")
    add_figure(d, "scene-08-dashboard.png", "Figure 5. Lecturer dashboard. The displayed cohort values come from the bundled simulated sample data.")
    d.add_heading("4.2 Add students and assessment marks", level=2)
    reset_steps(d)
    number(d, "Create or select a module and define its assessment plan, including weight, number of items, and how many items count where best-N-of-M is used.")
    number(d, "Add one student or import a supported student/mark CSV. The sample module and any demonstration data must remain simulated or anonymised.")
    number(d, "Enter marks on the 0–100 scale. Leave an uncompleted assessment blank. Review the validation messages if a mark is outside the accepted range.")
    d.add_heading("4.3 Review the at-risk list", level=2)
    d.add_paragraph("The at-risk view lists students whose maximum feasible final participation mark is below the module pass threshold. If no rows appear, the current data contains no student meeting that rule. This does not mean the system evaluated all forms of academic risk.")
    add_figure(d, "scene-08-at-risk.png", "Figure 6. At-risk view. In the bundled sample module, no student is currently flagged by the rule.")
    d.add_heading("4.4 Open a student and plan a target", level=2)
    reset_steps(d)
    number(d, "Select a student in the cohort list to open the individual detail view.")
    number(d, "Review the accrued participation mark, current standing, feasible final range and remaining assessment items.")
    number(d, "Choose an absolute target inside the displayed feasible range. The page shows the minimum uniform level required across remaining items and any skippable-item notes that apply.")
    d.add_paragraph("The lecturer DSS calculates and displays a plan for the lecturer. It does not send this plan to the student or create a notification.")
    add_figure(d, "scene-09-student-plan.png", "Figure 7. Lecturer individual plan with a feasible range and target requirements.")
    d.add_heading("4.5 Explore cohort assessment weights", level=2)
    reset_steps(d)
    number(d, "Open cohort planning and enter a target class average.")
    number(d, "For each assessment, enter a minimum and maximum weight. Minimums must be positive, the minimum cannot exceed the maximum, the ranges must allow a 100% total, and the solver enforces a 1% minimum weight.")
    number(d, "Calculate weights and review the solver status, proposed weights and resulting average. Treat this as a planning scenario rather than a guaranteed outcome.")
    add_figure(d, "scene-10-cohort-plan.png", "Figure 8. Lecturer cohort planner displays weight ranges and an optimisation result.")
    d.add_heading("5. Understand the results", level=1)
    table(d, ["Label", "How to read it"], [
        ("Participation so far", "Weighted contribution accrued from marks that currently count."),
        ("Standing", "Average across the work that has counted so far; it may differ from the accrued contribution."),
        ("Feasible range", "The minimum and maximum final mark possible under the entered assessment plan."),
        ("At risk", "The best-case feasible final mark is below the configured pass threshold."),
        ("Required mark", "The modelled mark level needed to reach the selected target, assuming the entered assessment plan and marks are correct."),
        ("Cohort weight proposal", "An optimisation result under the entered ranges and target; it is not a prediction of student behaviour."),
    ], [1.7, 5.0])
    d.add_heading("6. Privacy and responsible use", level=1)
    for s in [
        "Keep the tools separate: the lecturer DSS and student assistant do not share records.",
        "The student assistant saves data in the current browser's local storage. Use a private device/profile and avoid a shared public computer when entering personal details.",
        "Lecturer module data is held in server memory and is cleared when the lecturer service restarts. Use CSV export only where appropriate and store exported files securely.",
        "For demonstrations, use invented names and marks. The bundled lecturer cohort is simulated.",
        "APFIS does not read LMS logins, submissions or attendance activity and does not send email, push notifications or in-app messages to students.",
        "Use the output to support discussion and planning. Confirm marks, weights and pass rules against the official module information.",
    ]:
        bullet(d, s)
    add_figure(d, "scene-07-privacy.png", "Figure 9. The two tools operate independently and do not share student data.")
    d.add_heading("7. Troubleshooting", level=1)
    table(d, ["Problem", "What to check"], [
        ("Landing page says an app is unavailable", "Confirm the relevant local server is running on port 5001 or 5002; restart it if needed."),
        ("Student phases do not show saved information", "Use the same browser and origin; check that local storage is available and that the model was saved in Phase 1."),
        ("Returning CSV is rejected", "Choose the CSV exported by APFIS and avoid changing its section headings or rows."),
        ("Marks are not accepted", "Enter a number from 0 to 100; leave an unwritten assessment blank."),
        ("Lecturer data disappeared", "The lecturer store is in memory and resets when that service restarts. Re-import a suitable CSV or recreate the demo data."),
        ("No student appears on the at-risk page", "The rule only flags a student when their maximum feasible final mark is below the configured pass threshold."),
        ("Weight calculation reports invalid ranges", "Check positive minimums, minimum ≤ maximum, and whether all ranges can total exactly 100%."),
    ], [2.1, 4.6])
    d.add_paragraph("When asking a developer for help, include which tool and screen you were using, the steps immediately before the issue, and the exact validation message. Remove names, student numbers and marks from screenshots unless the data is fictional and approved for sharing.")
    d.add_heading("8. Quick workflow reference", level=1)
    table(d, ["Student assistant", "Lecturer DSS"], [
        ("Welcome → assessment plan → enter existing marks → generate profile → explore improvement options → review/export summary", "Choose module → define assessment plan → add/import cohort → review dashboard and at-risk rule → open student target plan → explore cohort weight ranges"),
    ], [3.35, 3.35])
    d.save(OUT / "APFIS_User_Manual.docx")

if __name__ == "__main__":
    technical_manual()
    user_manual()
    print(OUT / "APFIS_Developers_Technical_Manual.docx")
    print(OUT / "APFIS_User_Manual.docx")
